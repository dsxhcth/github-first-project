from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

OUT_DIR = r"D:\dsh"
DOCX = os.path.join(OUT_DIR, "运动学精讲题思路_手机版.docx")

doc = Document()

# Page setup
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.7)
sec.bottom_margin = Cm(1.6)
sec.left_margin = Cm(1.8)
sec.right_margin = Cm(1.8)

# Base styles
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.paragraph_format.line_spacing = 1.35
normal.paragraph_format.space_after = Pt(4)

for style_name, size, color in [("Heading 1", 15, "1F4E79"), ("Heading 2", 12, "2E75B6")]:
    st = styles[style_name]
    st.font.name = "Microsoft YaHei"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.paragraph_format.space_before = Pt(14)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True


def set_cn_font(run, name="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def shade(paragraph, fill="EAF3FB"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def add_title(text, sub):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string("1F4E79")
    set_cn_font(r)
    p.paragraph_format.space_after = Pt(2)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(sub)
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor.from_string("666666")
    set_cn_font(r2)
    p2.paragraph_format.space_after = Pt(10)


def add_body(text, bold=False, size=10.5, color=None, indent=False):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.0)
    if indent:
        p.paragraph_format.left_indent = Cm(0.4)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    set_cn_font(r)
    return p


def add_formula(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string("1F4E79")
    set_cn_font(r, "Cambria Math")
    return p


def add_answer(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    shade(p, "FFF2CC")
    r = p.add_run("答案：")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string("C00000")
    set_cn_font(r)
    r2 = p.add_run(text)
    r2.bold = True
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor.from_string("C00000")
    set_cn_font(r2)
    return p


def add_tip(num, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.first_line_indent = Cm(-0.3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"{num}. ")
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string("2E75B6")
    set_cn_font(r)
    r2 = p.add_run(text)
    r2.font.size = Pt(10.5)
    set_cn_font(r2)
    return p


def add_bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.55 + level * 0.4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    set_cn_font(r)
    return p


def add_page_number_footer():
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("运动学精讲题思路 · 手机阅读版 · 第 ")
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string("888888")
    set_cn_font(r)
    fld1 = OxmlElement("w:fldSimple")
    fld1.set(qn("w:instr"), "PAGE")
    run_el = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "17"); rPr.append(sz)
    color = OxmlElement("w:color"); color.set(qn("w:val"), "888888"); rPr.append(color)
    run_el.append(rPr)
    t = OxmlElement("w:t"); t.text = "1"; run_el.append(t)
    fld1.append(run_el)
    p._p.append(fld1)
    r2 = p.add_run(" 页")
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RGBColor.from_string("888888")
    set_cn_font(r2)


add_title("运动学精讲题思路", "手机阅读版｜4 道核心题 + 1 道挑战题｜家教专用")

p = doc.add_paragraph()
shade(p, "EAF3FB")
r = p.add_run("使用说明：每道题先看“思路”，再按“讲解技巧”提问。不要一次讲完，让学生先说下一步。")
r.font.size = Pt(9.5)
r.font.color.rgb = RGBColor.from_string("1F4E79")
set_cn_font(r)

# Quick box
doc.add_heading("先记住四句话", level=1)
for line in [
    "自由落体：先算总时间，再算某一段。",
    "v-t 图：斜率看加速度，面积看位移，t 轴下方取负。",
    "运动过程题：先口述“速度怎么变”，再选图。",
    "x/t-t 图：纵轴不是速度；斜率是 a/2，截距是 v₀。",
]:
    add_bullet(line)

# 1
doc.add_heading("一、P8 第10题：自由落体（80 m）", level=1)
add_body("题目核心：H = 80 m，g = 10 m/s²，自由落体。", bold=True)
add_body("解题思路：先求总时间。")
add_formula("t = √(2H/g) = √(2×80/10) = 4 s")
add_body("再逐项判断：")
add_bullet("A：第1 s内位移 h₁ = ½gt² = ½×10×1² = 5 m，不是10 m，错。")
add_bullet("B：前2 s位移 h₂ = ½×10×2² = 20 m；第2 s内位移 = 20 − 5 = 15 m，比第1 s多10 m，不是5 m，错。")
add_bullet("C：第1 s末速度 v₁ = 10 m/s，第2 s末速度 v₂ = 20 m/s，差10 m/s，对。")
add_bullet("D：下落总时间4 s，不是8 s，错。")
add_answer("C")
add_body("讲解技巧：", bold=True, color="2E75B6")
for i, t in enumerate([
    "先画时间轴，标出 0、1、2、4 s。",
    "反复问：“第2 s内”是哪一段？——是1 s到2 s，不是前2 s。",
    "教固定动作：先算累计位移，再相减。第2 s内 = h(2) − h(1)。",
    "学生选B时不要直接说错，让他自己算 h(1) 和 h(2)，通常算完就发现“多5 m”和“多10 m”弄反了。",
    "变式：把 H 改成45 m，让他再算 t = 3 s、落地速度30 m/s。",
], 1):
    add_tip(i, t)

# 2
doc.add_heading("二、P5 变式3-1：v-t 图像", level=1)
add_body("题目核心：v-t 图，斜率是加速度，面积是位移。", bold=True)
add_body("解题思路：")
add_bullet("A：第3 s内是2 s→3 s，速度从2 m/s降到0，a = (0−2)/1 = −2 m/s²，大小2，方向向西。A说向东，错。")
add_formula("s₁ = ½×3×4 = 6 m")
add_formula("s₂ = ½×2×(−2) = −2 m")
add_formula("总位移 = 6 − 2 = 4 m，平均速度 = 4/5 = 0.8 m/s，方向向东")
add_bullet("B：上面计算正确，对。")
add_bullet("C：t = 3 s 时速度由正变负，运动方向改变；不是 t = 4 s，错。")
add_bullet("D：1—3 s 斜率 −2，3—4 s 斜率也是 −2，方向相同，不是相反，错。")
add_answer("B")
add_body("讲解技巧：", bold=True, color="2E75B6")
for i, t in enumerate([
    "先让学生用手比划图线：“上坡还是下坡？”——上坡 a>0，下坡 a<0。",
    "面积只教三角形：½×底×高。先不算符号，算完再判断在 t 轴上方还是下方。",
    "强调三句话：斜率看加速度；面积看位移；t 轴下方面积取负。",
    "学生最容易把“第3 s内”当成“0—3 s”，让他用笔把2—3 s涂出来。",
    "讲完B后追问：“0—5 s的路程是多少？”——应是6+2=8 m，不是4 m。这个对比能帮他区分位移和路程。",
], 1):
    add_tip(i, t)

# 3
doc.add_heading("三、P7 第4题：运动过程→图像", level=1)
add_body("题目核心：小球从斜面静止滚下，进入水平面后匀减速到停止，选速率-时间图像。", bold=True)
add_body("解题思路：把运动拆成两段。")
add_bullet("斜面：从静止开始，速度均匀增大——图像从原点出发，斜率为正。")
add_bullet("水平面：匀减速到停止，速度均匀减小到0——图像下降，最后回到 t 轴。")
add_bullet("速率始终不小于0，所以图像不能穿到 t 轴下方。")
add_bullet("A：先减小后不变，不符合。")
add_bullet("B：先增大后不变，没有停止，不符合。")
add_bullet("C：先增大后减小到0，符合。")
add_bullet("D：先减小到0再增大，不符合。")
add_answer("C")
add_body("讲解技巧：", bold=True, color="2E75B6")
for i, t in enumerate([
    "不要一上来讲图像，先让他口述：“小球在斜面上速度怎么变？到水平面后怎么变？”",
    "用“速度表”比喻：斜面踩油门，水平面踩刹车，最后车停。",
    "强调“速率”与“速度”的区别：速率图永远在 t 轴上方；速度图才可能为负。",
    "排除法最省时间：先看起点是不是0，再看最后是不是0，能排除A、B、D。",
    "如果学生基础更弱，让他在草稿纸上画三个小场景：斜面、水平面、停止，再对应图像。",
], 1):
    add_tip(i, t)

# 4
doc.add_heading("四、P7 第5题：x/t-t 图像（易错）", level=1)
add_body("题目核心：纵轴是 x/t，不是速度，也不是位移。", bold=True)
add_body("解题思路：由位移公式两边除以 t。")
add_formula("x = v₀t + ½at²  →  x/t = v₀ + ½at")
add_body("所以 x/t-t 图像是一条直线：截距是 v₀，斜率是 ½a。")
add_body("图中直线过 (0,2) 和 (10,10)：")
add_formula("v₀ = 2 m/s")
add_formula("k = (10−2)/10 = 0.8 = ½a  →  a = 1.6 m/s²")
add_body("逐项判断：")
add_bullet("A：x(10) = 2×10 + ½×1.6×100 = 100 m，不是55 m，错。")
add_bullet("B：a = 1.6 m/s²，不是0.80，错。")
add_bullet("C：x(5) = 2×5 + ½×1.6×25 = 10 + 20 = 30 m，对。")
add_bullet("D：v(10) = v₀ + at = 2 + 1.6×10 = 18 m/s，不是10，错。")
add_answer("C")
add_body("讲解技巧：", bold=True, color="2E75B6")
for i, t in enumerate([
    "先问：“纵轴写的是什么？”——x/t，不是 v。这一步必须让学生自己说出来。",
    "把 x/t = v₀ + ½at 与 y = b + kx 对应，标出截距和斜率。",
    "强调：斜率不是加速度，斜率是加速度的一半。",
    "这题最容易错在把纵轴当速度，直接选D。可以让他先算 v(10)=18，再回头看D，印象会深。",
    "变式：把图线改成过 (0,4)、(10,14)，让他再求 v₀、a、x(10)。",
], 1):
    add_tip(i, t)

# 5
doc.add_heading("五、挑战题：P8 第6题（学有余力再讲）", level=1)
add_body("初速度为零的匀加速运动，第 n 秒内位移：sₙ = ½a(2n−1)。", bold=True)
add_formula("s₂ = 3a/2，s₃ = 5a/2")
add_body("第2 s内位移的前1/3是 a/2。设从第2 s开始用时 t₁，初速度 v₁ = a：")
add_formula("a/2 = a·t₁ + ½a·t₁²  →  t₁ = √2 − 1")
add_body("第3 s内位移的后1/5是 a/2。倒着算，末速度 v₃ = 3a：")
add_formula("a/2 = 3a·t₂ − ½a·t₂²  →  t₂ = 3 − 2√2 = (√2 − 1)²")
add_formula("t₂/t₁ = √2 − 1 ≈ 0.414")
add_body("在 1/3 和 1/2 之间。")
add_answer("A")
add_body("讲解技巧：这题对23分学生偏难，建议放在最后“挑战题”，不要求独立完成。只教他画时间轴、标 s₂、s₃，并记住“第n秒内位移比是1∶3∶5”。不要展开复杂运算。", color="666666")

# 6
doc.add_heading("六、明天上课的精讲顺序", level=1)
for i, t in enumerate([
    "P8 第10题：练自由落体，建立“先算总时间、再分段”的习惯。",
    "P5 变式3-1：练 v-t 图的斜率和面积。",
    "P7 第4题：练“运动过程→图像”的对应。",
    "P7 第5题：练 x/t-t 图像，防止把斜率当加速度。",
    "P8 第6题只作为挑战，不讲透。",
], 1):
    add_tip(i, t)
p = doc.add_paragraph()
shade(p, "EAF3FB")
r = p.add_run("每道题都坚持三句话：已知什么？要求什么？选哪条公式 / 看斜率还是面积？学生说对第一步，再让他动笔。")
r.font.size = Pt(9.5)
r.font.color.rgb = RGBColor.from_string("1F4E79")
set_cn_font(r)

add_page_number_footer()
doc.save(DOCX)
print("saved", DOCX)
